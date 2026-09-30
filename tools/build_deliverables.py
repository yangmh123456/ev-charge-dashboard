from __future__ import annotations

import json
import shutil
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PptInches
from pptx.util import Pt as PptPt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "deliverables"
SCREENSHOT = ROOT / "dashboard-screenshot.png"
sys.path.insert(0, str(ROOT))

from app.analytics import EVDataRepository, build_dashboard_payload, build_mapreduce_jobs, build_requirement_coverage


def ensure_dirs() -> None:
    OUT.mkdir(exist_ok=True)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = "Microsoft YaHei"


def add_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(21)
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.add_run(text)
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(11)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value


def build_report(payload: dict) -> Path:
    coverage = build_requirement_coverage()
    mapreduce_jobs = build_mapreduce_jobs(EVDataRepository())
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("新能源充电桩大数据分析平台实践报告")
    run.bold = True
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(20)
    doc.add_paragraph("项目名称：电力&新能源充电池桩大数据项目").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("技术路线：Flask + Pandas + SVG可视化 + MapReduce分析口径").alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(doc, "1 绪论", 1)
    add_paragraph(
        doc,
        "本项目面向新能源汽车充电桩运营场景，围绕充电会话、电池SOC、充电时长、费用、平台选择等数据进行统计分析和可视化展示。系统目标是让运营人员能够快速了解充电行为分布、设备运行状态和预测类指标，为站点运营管理提供数据支撑。",
    )

    add_heading(doc, "2 数据集", 1)
    summary = payload["summary"]
    add_table(
        doc,
        ["数据表", "来源文件", "记录数", "主要字段", "用途"],
        [
            ["充电会话数据", "nvv2t_md.csv", str(summary["session_count"]), "kwhTotal、chargeTimeHrs、platform、stationId", "充电频率、时长、费用、平台预测"],
            ["电池时序数据", "dsv13r2.csv", str(summary["battery_points"]), "record_time、soc、pack_voltage、temperature", "SOC均值、性能指标、温度分析"],
        ],
    )
    add_paragraph(
        doc,
        f"经程序读取和清洗后，本项目共使用 {summary['session_count']} 条充电会话记录、{summary['battery_points']} 条电池采样记录，累计充电量 {summary['total_kwh']} kWh，平均充电时长 {summary['avg_charge_hours']} 小时。",
    )

    add_heading(doc, "3 系统设计", 1)
    add_paragraph(
        doc,
        "系统采用分层结构：数据层读取课程CSV数据并完成类型清洗；分析层封装统计聚合、指标计算和简单预测逻辑；服务层使用Flask提供首页和JSON API；前端层使用HTML、CSS和原生SVG图表展示分析结果。这样的结构便于后续把本地Pandas分析替换为Hadoop/MapReduce离线结果。",
    )
    add_table(
        doc,
        ["模块", "文件", "说明"],
        [
            ["数据分析模块", "app/analytics.py", "负责CSV读取、清洗、聚合统计和预测"],
            ["Web服务模块", "app/web.py", "提供 / 首页、/api/dashboard 接口和健康检查"],
            ["前端展示模块", "static/index.html、styles.css、app.js", "展示项目进度、指标卡、图表、报告和预测结果"],
            ["自动化测试", "tests/", "验证数据读取、统计逻辑、API和首页"],
        ],
    )

    add_heading(doc, "4 功能实现", 1)
    rows = []
    for item in payload["sections"]:
        rows.append([item["id"], item["title"], item["type"]])
    add_table(doc, ["功能编号", "功能名称", "实现类型"], rows)
    add_paragraph(
        doc,
        "系统已经按照任务书要求接入10个功能入口，包括BI商业智能数据可视化、充电时间与次数频率分析、每时电池平均SOC、充电时间分布图分析、平均充电速率图分析、性能指标分类分析报告，以及四类预测功能。",
    )

    add_heading(doc, "5 运行结果", 1)
    add_paragraph(
        doc,
        "运行 Flask 服务后，浏览器访问 http://127.0.0.1:5000 可进入项目驾驶舱。首页展示项目进度、核心指标、统计图表、性能报告和历史数据估算结果。",
    )
    if SCREENSHOT.exists():
        doc.add_picture(str(SCREENSHOT), width=Inches(6.2))

    add_heading(doc, "6 Hadoop/MapReduce对应说明", 1)
    add_paragraph(
        doc,
        "课程要求中的Hadoop/HDFS/MapReduce主要承担大批量数据的离线存储和聚合计算。本项目当前在本机使用Pandas完成同等分析口径，后续可将 charging_frequency_by_hour、hourly_soc、time_distribution、average_speed_by_hour 等函数替换为MapReduce作业输出，再由Flask读取结果文件进行展示。",
    )
    add_table(
        doc,
        ["分析任务", "Map阶段", "Reduce阶段"],
        [
            ["充电时间频率", "按startTime输出(hour, 1, kWh)", "按hour汇总次数和电量"],
            ["每时平均SOC", "按record_time小时输出(hour, soc)", "按hour求SOC平均值"],
            ["充电时间分布", "按chargeTimeHrs映射区间", "按区间统计会话数"],
            ["平均充电速率", "输出(hour, kWh/time)", "按hour求平均速率"],
        ],
    )

    add_heading(doc, "7 测试与验证", 1)
    add_paragraph(
        doc,
        "项目提供pytest自动化测试，覆盖数据集读取、统计函数、预测函数、Flask API和首页。前端脚本通过node --check进行语法验证，浏览器页面通过Playwright截图检查渲染效果。",
    )

    add_heading(doc, "8 课程要求覆盖矩阵", 1)
    add_table(
        doc,
        ["要求项", "状态", "证据"],
        [[item["name"], item["status"], item["evidence"]] for item in coverage["items"]],
    )

    add_heading(doc, "9 MapReduce作业矩阵", 1)
    add_table(
        doc,
        ["作业", "Map键", "Reduce输出", "输出行数"],
        [[job["title"], job["map_key"], job["reduce_value"], str(len(job["rows"]))] for job in mapreduce_jobs["jobs"]],
    )

    add_heading(doc, "10 安全加固", 1)
    add_paragraph(
        doc,
        "系统已按企业级Web应用的基本安全要求进行加固：Flask debug默认关闭；统一添加CSP、X-Frame-Options、X-Content-Type-Options、Referrer-Policy等响应头；限制请求体大小；前端使用DOM API渲染文本和SVG，避免innerHTML类字符串注入风险。",
    )

    add_heading(doc, "11 总结", 1)
    add_paragraph(
        doc,
        "本项目完成了从课程数据集读取、数据分析、Web API、前端可视化到预测结果展示的闭环。通过模块化设计，系统既能满足当前课程演示和报告提交，也保留了向Hadoop/MapReduce离线计算扩展的接口。",
    )

    path = OUT / "新能源充电桩大数据分析实践报告.docx"
    doc.save(path)
    return path


def add_title(slide, text: str, color=RGBColor(238, 244, 234)) -> None:
    box = slide.shapes.add_textbox(PptInches(0.55), PptInches(0.35), PptInches(12.2), PptInches(0.65))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = PptPt(30)
    run.font.bold = True
    run.font.color.rgb = color


def add_body(slide, lines: list[str], x=0.75, y=1.25, w=5.6, h=4.8, size=17) -> None:
    box = slide.shapes.add_textbox(PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    frame = box.text_frame
    frame.clear()
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Microsoft YaHei"
        paragraph.font.size = PptPt(size)
        paragraph.font.color.rgb = RGBColor(238, 244, 234)
        paragraph.space_after = PptPt(8)


def add_stat(slide, x: float, y: float, label: str, value: str, accent=RGBColor(123, 216, 143)) -> None:
    shape = slide.shapes.add_shape(1, PptInches(x), PptInches(y), PptInches(2.55), PptInches(1.15))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(32, 39, 32)
    shape.line.color.rgb = RGBColor(64, 78, 64)
    box = slide.shapes.add_textbox(PptInches(x + 0.15), PptInches(y + 0.12), PptInches(2.25), PptInches(0.9))
    frame = box.text_frame
    frame.clear()
    p1 = frame.paragraphs[0]
    p1.text = label
    p1.font.name = "Microsoft YaHei"
    p1.font.size = PptPt(11)
    p1.font.color.rgb = RGBColor(154, 168, 154)
    p2 = frame.add_paragraph()
    p2.text = value
    p2.font.name = "Microsoft YaHei"
    p2.font.size = PptPt(23)
    p2.font.bold = True
    p2.font.color.rgb = accent


def set_dark_bg(slide) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(16, 19, 16)


def build_ppt(payload: dict) -> Path:
    coverage = build_requirement_coverage()
    mapreduce_jobs = build_mapreduce_jobs(EVDataRepository())
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)
    blank = prs.slide_layouts[6]
    summary = payload["summary"]

    slide = prs.slides.add_slide(blank)
    set_dark_bg(slide)
    add_title(slide, "新能源充电桩大数据分析平台")
    add_body(slide, ["课程项目答辩", "Flask + Pandas + 可视化驾驶舱", "面向充电行为统计、性能分析与预测展示"], y=1.5, w=7.2, size=22)
    add_stat(slide, 8.5, 1.5, "充电会话", str(summary["session_count"]))
    add_stat(slide, 8.5, 2.95, "累计电量", f"{summary['total_kwh']} kWh", RGBColor(241, 200, 75))
    add_stat(slide, 8.5, 4.4, "功能入口", "10")

    slide = prs.slides.add_slide(blank)
    set_dark_bg(slide)
    add_title(slide, "项目背景与目标")
    add_body(slide, ["背景：新能源汽车充电站需要从会话、电池和费用数据中发现运营规律。", "目标：实现数据读取、统计分析、预测展示和可视化驾驶舱。", "价值：辅助运营人员观察高峰时段、充电效率、平台偏好和设备状态。"], w=11.5)

    slide = prs.slides.add_slide(blank)
    set_dark_bg(slide)
    add_title(slide, "数据集概览")
    add_stat(slide, 0.8, 1.35, "会话数据", f"{summary['session_count']} 条")
    add_stat(slide, 3.65, 1.35, "电池采样", f"{summary['battery_points']} 条", RGBColor(115, 214, 208))
    add_stat(slide, 6.5, 1.35, "平台数量", str(summary["platforms"]), RGBColor(241, 200, 75))
    add_stat(slide, 9.35, 1.35, "平均时长", f"{summary['avg_charge_hours']} h")
    add_body(slide, ["nvv2t_md.csv：充电会话、时长、电量、费用、平台、站点。", "dsv13r2.csv：电池记录时间、SOC、电压、电流、温度、可用容量。"], y=3.2, w=11.6)

    slide = prs.slides.add_slide(blank)
    set_dark_bg(slide)
    add_title(slide, "系统架构")
    add_body(slide, ["数据层：读取课程CSV并完成字段类型清洗。", "分析层：封装统计聚合、性能指标和预测逻辑。", "服务层：Flask提供首页与 /api/dashboard。", "前端层：HTML/CSS/JS + SVG图表展示数据。"], w=6.1)
    if SCREENSHOT.exists():
        slide.shapes.add_picture(str(SCREENSHOT), PptInches(7.0), PptInches(1.25), width=PptInches(5.6))

    slide = prs.slides.add_slide(blank)
    set_dark_bg(slide)
    add_title(slide, "核心功能清单")
    add_body(slide, [f"{index + 1}. {item['title']}" for index, item in enumerate(payload["sections"][:10])], w=11.8, size=15)

    slide = prs.slides.add_slide(blank)
    set_dark_bg(slide)
    add_title(slide, "可视化运行结果")
    if SCREENSHOT.exists():
        slide.shapes.add_picture(str(SCREENSHOT), PptInches(0.7), PptInches(1.15), width=PptInches(11.9))

    slide = prs.slides.add_slide(blank)
    set_dark_bg(slide)
    add_title(slide, "Hadoop / MapReduce 对应关系")
    add_body(slide, ["充电频率：Map输出小时键值，Reduce汇总次数与电量。", "每时SOC：Map提取记录小时与SOC，Reduce求平均值。", "时间分布：Map映射时长区间，Reduce统计区间数量。", "平均速率：Map计算单次速率，Reduce按小时求均值。", "当前本机版使用Pandas实现同等分析口径，后续可替换为HDFS离线结果。"], w=11.7, size=16)

    slide = prs.slides.add_slide(blank)
    set_dark_bg(slide)
    add_title(slide, "课程要求覆盖矩阵")
    add_body(slide, [f"{item['name']}：{item['status']}" for item in coverage["items"]], w=11.7, size=17)

    slide = prs.slides.add_slide(blank)
    set_dark_bg(slide)
    add_title(slide, "MapReduce 作业矩阵")
    add_body(slide, [f"{job['title']}：{job['map_key']} -> {job['reduce_value']}，输出{len(job['rows'])}行" for job in mapreduce_jobs["jobs"]], w=11.7, size=15)

    slide = prs.slides.add_slide(blank)
    set_dark_bg(slide)
    add_title(slide, "安全加固")
    add_body(slide, ["Flask debug 默认关闭，按环境变量显式开启。", "统一设置 CSP、X-Frame-Options、X-Content-Type-Options 等安全响应头。", "限制请求体大小，保留 HTTPS Cookie 配置开关。", "前端删除 innerHTML 渲染，改用 textContent 与 DOM API。"], w=11.7, size=17)

    slide = prs.slides.add_slide(blank)
    set_dark_bg(slide)
    add_title(slide, "测试与验证")
    add_body(slide, ["pytest tests -q：验证数据读取、统计函数、预测函数、API和首页。", "node --check static/app.js：验证前端脚本语法。", "Invoke-RestMethod /api/dashboard：验证API返回真实汇总数据。", "Playwright screenshot：验证浏览器页面正常渲染。"], w=11.7, size=17)

    slide = prs.slides.add_slide(blank)
    set_dark_bg(slide)
    add_title(slide, "项目总结")
    add_body(slide, ["完成了从数据读取到可视化展示的闭环。", "覆盖课程要求中的10个功能入口。", "系统结构清晰，便于扩展到Hadoop/MapReduce离线计算。", "交付物包含源码、报告、PPT、截图和运行说明。"], w=11.6, size=19)

    path = OUT / "新能源充电桩大数据分析答辩PPT.pptx"
    prs.save(path)
    return path


def build_mapreduce_note() -> Path:
    text = """# Hadoop/MapReduce 补充说明

本项目当前在 Windows 本机用 Flask + Pandas 实现课程数据的读取、分析和可视化。为了对应课程要求中的 Hadoop、HDFS、MapReduce，系统设计时把统计分析逻辑集中在 `app/analytics.py`，这些函数可以一一替换为 MapReduce 离线作业输出。

## HDFS 目录规划

```text
/ev_charge/input/nvv2t_md.csv
/ev_charge/input/dsv13r2.csv
/ev_charge/output/time_frequency
/ev_charge/output/hourly_soc
/ev_charge/output/time_distribution
/ev_charge/output/avg_speed
```

## MapReduce 任务口径

| 功能 | Map 输出 | Reduce 输出 |
| --- | --- | --- |
| 充电时间与次数频率分析 | `(startTime, 1, kwhTotal)` | 每小时会话数、总充电量 |
| 每时电池平均SOC | `(record_hour, soc)` | 每小时SOC平均值 |
| 充电时间分布图分析 | `(duration_bucket, 1)` | 每个时长区间的会话数 |
| 平均充电速率图分析 | `(startTime, kwhTotal / chargeTimeHrs)` | 每小时平均充电速率 |

## 与当前代码的对应关系

- `charging_frequency_by_hour()` 对应充电频率 MapReduce 作业。
- `hourly_soc()` 对应 SOC 平均值 MapReduce 作业。
- `time_distribution()` 对应充电时长分布 MapReduce 作业。
- `average_speed_by_hour()` 对应平均速率 MapReduce 作业。

Flask 前端不关心结果来自 Pandas 还是 Hadoop，只要保持 `/api/dashboard` 的 JSON 结构一致即可。
"""
    path = OUT / "Hadoop_MapReduce补充说明.md"
    path.write_text(text, encoding="utf-8")
    return path


def build_submission_readme(paths: list[Path]) -> Path:
    payload_path = OUT / "提交说明.txt"
    payload_path.write_text(
        "新能源充电桩大数据分析课程项目提交说明\n\n"
        "运行方式：\n"
        "1. 进入项目目录 <当前仓库目录>\n"
        "2. 执行 python -m pip install -r requirements.txt\n"
        "3. 执行 python -m app.web\n"
        "4. 浏览器打开 http://127.0.0.1:5000\n\n"
        "交付物：\n"
        + "\n".join(f"- {path.name}" for path in paths)
        + "\n\n源码目录包含 app、static、tests、README.md、requirements.txt。\n",
        encoding="utf-8",
    )
    return payload_path


def build_zip(extra_paths: list[Path]) -> Path:
    zip_path = OUT / "新能源充电桩大数据分析课程项目提交包.zip"
    include_roots = [ROOT / "app", ROOT / "static", ROOT / "tests", ROOT / "tools"]
    include_files = [ROOT / "README.md", ROOT / "requirements.txt", ROOT / "dashboard-screenshot.png", ROOT / "security_best_practices_report.md", *extra_paths]
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for folder in include_roots:
            for path in folder.rglob("*"):
                if "__pycache__" in path.parts:
                    continue
                archive.write(path, path.relative_to(ROOT))
        for path in include_files:
            if path.exists() and path != zip_path:
                archive.write(path, path.relative_to(ROOT))
    return zip_path


def main() -> None:
    ensure_dirs()
    payload = build_dashboard_payload(EVDataRepository())
    report = build_report(payload)
    ppt = build_ppt(payload)
    note = build_mapreduce_note()
    security = ROOT / "security_best_practices_report.md"
    readme = build_submission_readme([report, ppt, note, security, SCREENSHOT])
    package = build_zip([report, ppt, note, readme])
    print(json.dumps({"report": str(report), "ppt": str(ppt), "note": str(note), "readme": str(readme), "zip": str(package)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
